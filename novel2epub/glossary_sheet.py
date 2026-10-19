"""Bảng tính Excel (.xlsx) cho glossary: dựng file xuất và đọc ngược để nhập.

Tách khỏi `app/routes` để test được không cần HTTP và để mọi đường xuất/nhập
dùng chung MỘT bộ tên cột — đổi ở đây thì cả file xuất lẫn cách nhận diện
cột khi người dùng sửa tên trong Excel đều đi theo.

Định dạng: 1 sheet, dòng tiêu đề cố định rồi mỗi mục 1 dòng. Ba cột đầu là
dữ liệu thật (Hán / Việt / Ghi chú); các cột sau (`Version`, `Cập nhật`) chỉ để
người dùng đối chiếu với bảng trên web và BỎ QUA khi nhập.
"""
from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO

from .glossary_review import INVALID_SOURCE_DETAIL
from .han_cleanup import count_han

try:  # dependency tuỳ chọn — thiếu thì route báo lỗi rõ thay vì 500 im lặng
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - chỉ xảy ra khi chưa cài openpyxl
    Workbook = None  # type: ignore[assignment]
    load_workbook = None  # type: ignore[assignment]

MISSING_OPENPYXL = (
    "Chưa cài openpyxl nên không xuất/nhập được file .xlsx. "
    "Chạy: pip install openpyxl"
)

SHEET_TITLE = "Glossary"

#: (tiêu đề, độ rộng cột) — thứ tự là thứ tự cột trong file.
COLUMNS: tuple[tuple[str, int], ...] = (
    ("Hán", 22),
    ("Việt", 24),
    ("Ghi chú", 40),
    ("Version", 9),
    ("Cập nhật", 18),
)

#: Tên cột được nhận khi đọc file — gồm cả tên tiếng Anh để file tạo ở nơi
#: khác vẫn vào được. So khớp TUYỆT ĐỐI (không substring) sau khi bỏ dấu và
#: mọi ký tự không phải chữ/số: "Hán tự" và "Hán tự (chinese)" → `hantu`,
#: `hantuchinese` — cái sau không khớp nên người dùng sửa tên cột cho gọn.
_COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "source": ("han", "hantu", "hanchu", "chuhan", "chinese", "source"),
    "target": ("viet", "vietnamese", "bandich", "dich", "meaning", "target"),
    "note": ("ghichu", "note", "notes", "chuthich", "comment", "comments"),
}

_NORMALIZE_RE = re.compile(r"[^a-z0-9]+")


def _normalize(text: str) -> str:
    """Bỏ dấu tiếng Việt + hạ chữ thường + bỏ hết ký tự lạ (space, `_`, `-`…)."""
    plain = "".join(
        ch for ch in unicodedata.normalize("NFD", str(text)) if not unicodedata.combining(ch)
    )
    return _NORMALIZE_RE.sub("", plain.lower())


def _cell_text(value) -> str:
    """Ô Excel → chuỗi đã trim (số/float ngày thì bỏ phần `.0`)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y %H:%M")
    return str(value).strip()


@dataclass
class SheetError:
    """Một dòng trong file không dùng được — báo lại để người dùng sửa."""

    row: int
    message: str

    def as_dict(self) -> dict:
        return {"row": self.row, "message": self.message}


@dataclass
class ParsedSheet:
    """Kết quả đọc file: mục hợp lệ + các dòng bị loại kèm lý do."""

    entries: list[tuple[str, str, str]] = field(default_factory=list)
    errors: list[SheetError] = field(default_factory=list)


def _require_openpyxl() -> None:
    if Workbook is None or load_workbook is None:  # pragma: no cover
        raise RuntimeError(MISSING_OPENPYXL)


def build_glossary_xlsx(
    entries: list[tuple[str, str, str]],
    meta: dict[str, tuple[int, float]] | None = None,
) -> bytes:
    """Dựng file .xlsx từ mục `(source, target, note)`.

    `meta` là `{source: (version, updated_at)}` từ `Storage.read_glossary_meta`
    — chỉ để đối chiếu, có thể bỏ trống (hai cột cuối để trống).
    """
    _require_openpyxl()
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_TITLE

    for index, (_title, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.append([title for title, _w in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"

    meta = meta or {}
    for source, target, note in entries:
        source = (source or "").strip()
        target = (target or "").strip()
        if not source or not target:
            continue
        version, updated_at = meta.get(source, (None, None))
        when = datetime.fromtimestamp(updated_at) if updated_at else None
        ws.append([source, target, (note or "").strip(), version, when or ""])

    if ws.max_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{ws.max_row}"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _header_columns(row: tuple) -> tuple[int, int, int] | None:
    """Tìm (cột Hán, Việt, Ghi chú) theo tên tiêu đề. None nếu dòng đầu không
    phải tiêu đề → khi đó coi luôn là A/B/C."""
    wanted: dict[str, int] = {}
    for index, cell in enumerate(row):
        key = _normalize(_cell_text(cell))
        if not key:
            continue
        for field_name, aliases in _COLUMN_ALIASES.items():
            if key in aliases and field_name not in wanted:
                wanted[field_name] = index
    if "source" in wanted and "target" in wanted:
        return (
            wanted["source"],
            wanted["target"],
            wanted.get("note", wanted["target"] + 1),
        )
    return None


def parse_glossary_xlsx(data: bytes) -> ParsedSheet:
    """Đọc file .xlsx thành mục `(source, target, note)` + danh sách dòng lỗi.

    - Dòng tiêu đề được nhận diện theo tên (Hán/Việt/Ghi chú, cả tiếng Anh); nếu
      không khớp thì mặc định cột A/B/C — để file người dùng tự tạo vẫn vào được.
    - Dòng trống hoàn toàn bị bỏ qua; source trùng thì dòng SAU thắng (giống
      `write_glossary_entries`).
    - Dòng thiếu Hán/Việt hoặc source không có chữ Hán KHÔNG ném lỗi cả file:
      bỏ dòng đó và ghi vào `errors` để người dùng biết mình bỏ sót gì.
    """
    _require_openpyxl()
    parsed = ParsedSheet()
    merged: dict[str, tuple[str, str]] = {}
    order: list[str] = []
    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            return parsed
        cols = _header_columns(header) or (0, 1, 2)
        # Dòng tiêu đề chiếm 1 hàng trong toạ độ Excel (để báo lỗi cho đúng).
        for offset, row in enumerate(rows):
            excel_row = 2 + offset
            if not row or all(_cell_text(c) == "" for c in row):
                continue
            source = _cell_text(row[cols[0]] if cols[0] < len(row) else None)
            target = _cell_text(row[cols[1]] if cols[1] < len(row) else None)
            note = _cell_text(row[cols[2]] if cols[2] < len(row) else None)
            if not source and not target and not note:
                continue
            if not source:
                parsed.errors.append(SheetError(excel_row, "Thiếu Hán."))
                continue
            if not target:
                parsed.errors.append(SheetError(excel_row, "Thiếu Việt."))
                continue
            if count_han(source) == 0:
                parsed.errors.append(SheetError(excel_row, INVALID_SOURCE_DETAIL))
                continue
            if source not in merged:
                order.append(source)
            merged[source] = (target, note)
    finally:
        wb.close()

    parsed.entries = [(s, merged[s][0], merged[s][1]) for s in order]
    return parsed